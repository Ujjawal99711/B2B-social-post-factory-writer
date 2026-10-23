"""One Word doc: founder posts and company posts, each caption with its images."""
from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)


def add_meta(doc: Document, label: str, value: str) -> None:
    if not value:
        return
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    r.bold = True
    p.add_run(value)


def add_body(doc: Document, text: str) -> None:
    for block in (text or "").replace("\r\n", "\n").split("\n"):
        p = doc.add_paragraph(block if block.strip() else "")
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(11)


def add_image(doc: Document, path: Path, width_in: float = 4.8) -> None:
    if not path.exists():
        doc.add_paragraph(f"[Missing image: {path.name}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))


def resolve_image(run_dir: Path, rel: str) -> Path:
    posts_dir = run_dir / "04_posts"
    if not rel:
        return Path()
    candidate = (run_dir / rel).resolve()
    if candidate.exists():
        return candidate
    named = posts_dir / Path(rel).name
    return named


def load_posts(run_dir: Path) -> list[dict]:
    posts = []
    for path in sorted((run_dir / "04_posts").glob("post_*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        posts.append(data)
    return posts


def add_post(doc: Document, run_dir: Path, index: int, post: dict) -> None:
    fmt = post.get("format") or "Post"
    title = post.get("working_title") or ""
    add_heading(doc, f"{index}. {fmt} — {title}", level=2)
    add_meta(doc, "Day", str(post.get("day") or ""))
    add_meta(doc, "Pillar", str(post.get("pillar") or ""))
    if post.get("company_post_type"):
        add_meta(doc, "Post type", str(post.get("company_post_type")))
    add_heading(doc, "Post text", level=3)
    add_body(doc, post.get("post_text") or "")

    if fmt == "Infographic":
        add_heading(doc, "Infographic", level=3)
        add_image(doc, resolve_image(run_dir, post.get("image_path") or ""), width_in=5.0)

    if fmt == "Carousel":
        add_heading(doc, "Carousel", level=3)
        for slide in post.get("slides") or []:
            sn = slide.get("slide_number", "")
            role = slide.get("slide_role", "")
            headline = slide.get("headline", "")
            add_heading(doc, f"Slide {sn} ({role}): {headline}", level=4)
            if slide.get("body_text"):
                add_body(doc, slide.get("body_text") or "")
            add_image(doc, resolve_image(run_dir, slide.get("image_path") or ""), width_in=4.6)


def build(founder_dir: Path, company_dir: Path, founder_name: str, out_path: Path) -> None:
    founder_summary = json.loads((founder_dir / "00_summary.json").read_text(encoding="utf-8"))
    company_summary = json.loads((company_dir / "00_summary.json").read_text(encoding="utf-8"))

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    add_heading(doc, founder_name, 0)
    add_meta(doc, "Company", str(company_summary.get("company_title") or "Supportact Enterprise Pvt Ltd"))
    add_meta(doc, "Positioning", str(founder_summary.get("selected_positioning_mode") or ""))
    add_meta(doc, "Growth goal", str(founder_summary.get("selected_growth_goal") or ""))
    doc.add_paragraph()

    add_heading(doc, "Founder profile posts", level=1)
    add_meta(doc, "Headline", str(founder_summary.get("headline") or ""))
    add_meta(doc, "Run", founder_dir.name)
    doc.add_paragraph()
    founder_posts = load_posts(founder_dir)
    for i, post in enumerate(founder_posts, start=1):
        add_post(doc, founder_dir, i, post)
        doc.add_page_break()

    add_heading(doc, "Company page posts", level=1)
    add_meta(doc, "Title", str(company_summary.get("company_title") or ""))
    add_meta(doc, "Tagline", str(company_summary.get("tagline_recommended") or ""))
    add_meta(doc, "Run", company_dir.name)
    doc.add_paragraph()
    company_posts = load_posts(company_dir)
    for i, post in enumerate(company_posts, start=1):
        add_post(doc, company_dir, i, post)
        if i < len(company_posts):
            doc.add_page_break()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Wrote {out_path}")
    print(f"founder_posts={len(founder_posts)} company_posts={len(company_posts)}")


def main() -> int:
    founder = ROOT / "outputs" / "run_20261005_173236"
    company = ROOT / "outputs" / "run_20261005_174339"
    out = Path.home() / "Downloads" / "Srabanee_Chakraborty_LinkedIn_posts.docx"
    if len(sys.argv) > 1:
        out = Path(sys.argv[1])
    build(founder, company, "Srabanee Chakraborty", out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
