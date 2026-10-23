"""Build a Word doc of all posts + visuals for a pipeline run."""
from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)


def add_meta(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    r.bold = True
    p.add_run(value)


def add_body(doc: Document, text: str) -> None:
    for block in (text or "").strip().split("\n"):
        p = doc.add_paragraph(block if block.strip() else "")
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(11)


def add_image(doc: Document, path: Path, width_in: float = 5.5) -> None:
    if not path.exists():
        doc.add_paragraph(f"[Missing image: {path.name}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))


def load_posts(posts_dir: Path) -> list[dict]:
    posts = []
    for path in sorted(posts_dir.glob("post_*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        data["_file"] = path.name
        posts.append(data)
    return posts


def build_doc(run_dir: Path, out_path: Path) -> None:
    posts_dir = run_dir / "04_posts"
    summary_path = run_dir / "00_summary.json"
    lmo_path = run_dir / "02_lmo.json"
    summary = {}
    if summary_path.exists():
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
    lmo = {}
    if lmo_path.exists():
        lmo = json.loads(lmo_path.read_text(encoding="utf-8"))
    profile = lmo.get("profile_optimization") or {}

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    add_heading(doc, "LinkedIn Content Pack", 0)
    add_meta(doc, "Run", summary.get("run_id", run_dir.name))
    add_meta(
        doc,
        "Positioning",
        str(summary.get("selected_positioning_mode") or lmo.get("selected_positioning_mode") or ""),
    )
    add_meta(
        doc,
        "Growth goal",
        str(summary.get("selected_growth_goal") or lmo.get("selected_growth_goal") or ""),
    )
    add_meta(doc, "Niche", str(summary.get("confirmed_niche") or lmo.get("confirmed_niche") or ""))
    doc.add_paragraph()

    add_heading(doc, "LMO — Profile optimization", level=1)

    headline = ((profile.get("headline") or {}).get("fixed") or summary.get("headline") or "")
    add_heading(doc, "Headline", level=2)
    add_body(doc, str(headline))

    about = (profile.get("about_section") or {}).get("fixed") or ""
    add_heading(doc, "About", level=2)
    add_body(doc, str(about))

    experience = (profile.get("current_role_description") or {}).get("fixed") or ""
    add_heading(doc, "Experience (current role)", level=2)
    add_body(doc, str(experience))

    doc.add_page_break()
    add_heading(doc, "Weekly posts", level=1)

    posts = load_posts(posts_dir)
    for idx, post in enumerate(posts):
        fmt = post.get("format", "Post")
        day = post.get("day", "")
        title = post.get("working_title", "")
        n = idx + 1
        add_heading(doc, f"{n}. {fmt} — {title}", level=1)
        add_meta(doc, "Day", day)
        if post.get("pillar"):
            add_meta(doc, "Pillar", str(post.get("pillar")))
        if post.get("growth_angle"):
            add_meta(
                doc,
                "Growth tags",
                f"{post.get('growth_angle')} | {post.get('proof_type')} | {post.get('cta_style')}",
            )

        add_heading(doc, "Caption / Post text", level=2)
        add_body(doc, post.get("post_text", ""))

        if fmt == "Infographic":
            add_heading(doc, "Infographic visual", level=2)
            img = post.get("image_path") or ""
            img_path = (run_dir / img).resolve() if img else Path()
            if not img_path.exists() and img:
                img_path = posts_dir / Path(img).name
            if not img_path.exists():
                img_path = posts_dir / f"post_{n:02d}_infographic.png"
            add_image(doc, img_path, width_in=5.2)

        if fmt == "Carousel":
            add_heading(doc, "Carousel slides", level=2)
            for slide in post.get("slides") or []:
                sn = slide.get("slide_number", "")
                role = slide.get("slide_role", "")
                hl = slide.get("headline", "")
                body = slide.get("body_text", "")
                add_heading(doc, f"Slide {sn} ({role}): {hl}", level=3)
                if body:
                    add_body(doc, body)
                sp = slide.get("image_path") or ""
                slide_path = (run_dir / sp).resolve() if sp else Path()
                if not slide_path.exists() and sp:
                    slide_path = posts_dir / Path(sp).name
                add_image(doc, slide_path, width_in=4.8)

        if idx < len(posts) - 1:
            doc.add_page_break()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Wrote {out_path}")


def main() -> int:
    run_id = sys.argv[1] if len(sys.argv) > 1 else "run_20260921_181300"
    run_dir = ROOT / "outputs" / run_id
    if not run_dir.exists():
        # allow full path to run dir
        alt = Path(sys.argv[1]) if len(sys.argv) > 1 else Path()
        if alt.exists() and (alt / "04_posts").exists():
            run_dir = alt
            run_id = run_dir.name
        else:
            print(f"FAIL: missing {run_dir}", file=sys.stderr)
            return 1
    if len(sys.argv) > 2:
        out_path = Path(sys.argv[2])
    else:
        out_path = run_dir / f"{run_id}_content_pack.docx"
    build_doc(run_dir, out_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
