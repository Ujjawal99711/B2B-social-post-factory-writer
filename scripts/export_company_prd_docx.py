"""Export Company Page track PRD as a Word doc with full System / User / Schema per stage.

Usage:
  python scripts/export_company_prd_docx.py
  python scripts/export_company_prd_docx.py path/to/out.docx
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PROMPTS = ROOT / "prompts"
SCHEMAS = ROOT / "schemas"

STAGES = [
    {
        "num": "00",
        "title": "Ingest + Normalize",
        "model": "Local — scripts/ingest_normalize.py (no LLM)",
        "prompt": None,
        "schema": None,
        "user_prompt": (
            "Input path: legacy .txt file OR multi-upload pack folder.\n\n"
            "Example:\n"
            "  python scripts/run_pipeline.py inputs/manju_travfounders_raw.txt --track company_page\n"
            "  python scripts/run_pipeline.py inputs/packs/<customer> --track company_page"
        ),
        "notes": (
            "Writes 00_normalized_raw.txt and 00_ingest_manifest.json. "
            "Optional Gemini vision for pack images if GEMINI_API_KEY set."
        ),
    },
    {
        "num": "01",
        "title": "RAW Input Structurer (shared with founder track)",
        "model": "gpt-6-luna (default). Rollback: TEXT_PROVIDER=gemini",
        "prompt": "01_raw_input_structurer.md",
        "schema": "01_raw_output.schema.json",
        "user_prompt": (
            "Normalized raw text from Stage 00 (contents of 00_normalized_raw.txt).\n\n"
            "Customer paste only — not the system prompt.\n\n"
            "Example shape:\n"
            "=== ONBOARDING / POSITIONING ===\n"
            "selected_positioning_mode: founder_plus_product | founder_only\n"
            "selected_growth_goal: lead_generation | personal_branding\n"
            "declared_niche: ...\n"
            "target_audience: ...\n"
            "...\n"
            "=== LINKEDIN PROFILE ===\n"
            "...\n"
            "=== COMPANY PAGE ===\n"
            "...\n"
            "=== PAST POSTS ===\n"
            "..."
        ),
        "notes": (
            "Output: 01_raw.json + 01_agent_1_input.json. "
            "Optional business_data.website when URL present."
        ),
    },
    {
        "num": "02b",
        "title": "Company Page LMO",
        "model": "gpt-6-luna (default). Rollback: TEXT_PROVIDER=gemini",
        "prompt": "02_company_lmo.md",
        "schema": "02_company_lmo_output.schema.json",
        "user_prompt": (
            "Pretty-printed JSON: 01_agent_1_input.json\n\n"
            "Same payload as personal LMO. Do not re-send the RAW system prompt.\n\n"
            "Runner: python scripts/run_pipeline.py inputs/your_raw.txt --track company_page"
        ),
        "notes": (
            "Output: 02_company_lmo.json + 02_company_handoff.json. "
            "Official title passthrough, 5 taglines (≤120) + recommended, "
            "About 1,000–1,500 chars, Brand DNA, company pillars, evidence bank."
        ),
    },
    {
        "num": "03c",
        "title": "Company Content Matrix",
        "model": "gpt-6-luna (default). Rollback: TEXT_PROVIDER=gemini",
        "prompt": "03_company_content_matrix.md",
        "schema": "03_company_content_matrix_output.schema.json",
        "user_prompt": (
            "Pretty-printed JSON: 02_company_handoff.json\n\n"
            "Company handoff only (page_type=company, brand_dna, pillars, "
            "business_summary, evidence, voice_notes, content_mode_constraints)."
        ),
        "notes": (
            "Exactly 4 briefs: 2 Infographic + 2 Carousel. "
            "company_post_type: educational_howto | case_study | behind_the_scenes | company_news. "
            "No polls. No Text."
        ),
    },
    {
        "num": "05c",
        "title": "Company Infographic Writer + Image",
        "model": "gpt-6-luna → OpenAI gpt-image-2",
        "prompt": "05_company_infographic_post.md",
        "schema": "05_infographic_output.schema.json",
        "user_prompt": (
            "Built by build_writer_payload(company_handoff, post_brief):\n\n"
            "=== HANDOFF ===\n"
            "{ full 02_company_handoff.json }\n\n"
            "=== POST_BRIEF ===\n"
            "{ one Infographic object from 03_company_content_matrix.json }\n\n"
            "Image step: runner prepends 00_brand_dna_lock.txt to image_generation_prompt "
            "and calls gpt-image-2 (not part of this system/user pair)."
        ),
        "notes": (
            "Caption in post_text. PNG: 04_posts/post_XX_infographic.png. "
            "Company voice; company CTA; same Brand DNA lock as carousel."
        ),
    },
    {
        "num": "06c",
        "title": "Company Carousel Writer + Slide Images",
        "model": "gpt-6-luna → OpenAI gpt-image-2 (one PNG per slide)",
        "prompt": "06_company_carousel_post.md",
        "schema": "06_carousel_output.schema.json",
        "user_prompt": (
            "Built by build_writer_payload(company_handoff, post_brief):\n\n"
            "=== HANDOFF ===\n"
            "{ full 02_company_handoff.json }\n\n"
            "=== POST_BRIEF ===\n"
            "{ one Carousel object from 03_company_content_matrix.json }\n\n"
            "Image step: one gpt-image-2 call per slide with Brand DNA lock "
            "(hard cap 14 slides)."
        ),
        "notes": (
            "Caption in post_text. Slides: 04_posts/post_XX_carousel_slide_YY.png. "
            "Cover + Content + CTA; company CTA slide."
        ),
    },
]


def set_run_font(run, *, size: int = 10, mono: bool = False, bold: bool = False) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    if mono:
        run.font.name = "Consolas"
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), "Consolas")
        rFonts.set(qn("w:hAnsi"), "Consolas")
        rFonts.set(qn("w:cs"), "Consolas")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_body(doc: Document, text: str, *, size: int = 10) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size)


def add_meta(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10, bold=True)
    r2 = p.add_run(value)
    set_run_font(r2, size=10)


def add_code_block(doc: Document, text: str, *, size: int = 8) -> None:
    for line in text.splitlines() or [""]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.15)
        run = p.add_run(line if line else " ")
        set_run_font(run, size=size, mono=True)


def build_doc(out_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    add_heading(doc, "PRD — LinkedIn Company Page Process", 0)
    add_body(
        doc,
        "Product requirements + full prompt pack for the company_page track. "
        "Founder track is unchanged and separate.",
    )

    add_heading(doc, "1. Product summary", 1)
    add_meta(doc, "Track", "--track company_page")
    add_meta(doc, "Goal", "Company page title + taglines + About, then 4 company posts (2 Infographic + 2 Carousel) with Brand DNA–locked images")
    add_meta(doc, "Text/JSON model", "OpenAI gpt-6-luna (default). Rollback: TEXT_PROVIDER=gemini → gemini-3.5-flash")
    add_meta(doc, "Image model", "OpenAI gpt-image-2")
    add_meta(doc, "Runner", "scripts/run_pipeline.py")

    add_heading(doc, "CLI", 2)
    add_code_block(
        doc,
        "python scripts/run_pipeline.py inputs/<raw>.txt --track company_page\n"
        "python scripts/run_pipeline.py inputs/<raw>.txt --track company_page --stop-after company_lmo\n"
        "python scripts/run_pipeline.py inputs/<raw>.txt --track company_page --stop-after company_matrix",
        size=9,
    )

    add_heading(doc, "2. Architecture", 1)
    add_code_block(
        doc,
        "Upload pack OR legacy .txt\n"
        "  → 00 Ingest + Normalize          (local)\n"
        "  → 01 RAW Input Structurer        (gpt-6-luna)  [shared with founder]\n"
        "  → 02b Company Page LMO           (gpt-6-luna)  title + 5 taglines + About + handoff/Brand DNA\n"
        "  → 03c Company Content Matrix     (gpt-6-luna)  exactly 2 Infographic + 2 Carousel\n"
        "  → 05c Company Infographic        (gpt-6-luna + gpt-image-2)\n"
        "  → 06c Company Carousel           (gpt-6-luna + gpt-image-2 per slide)\n"
        "  → stop",
        size=9,
    )
    add_body(
        doc,
        "Call pattern: call_text_llm(system_prompt=full prompts/*.md, user_content=payload only).",
    )

    add_heading(doc, "3. Locked product rules", 1)
    add_body(
        doc,
        "• Voice: company we / company name — never founder I/me\n"
        "• Formats: exactly 2 Infographic + 2 Carousel (no Text, no polls)\n"
        "• Types: educational_howto | case_study | behind_the_scenes | company_news\n"
        "• At least one educational_howto; never invent case/news/BTS evidence\n"
        "• Tagline ≤ 120 characters; 5 options + 1 recommended pick\n"
        "• About: 1,000–1,500 characters; title = official name passthrough\n"
        "• CTA: visit_website | contact_team | follow_page | none (never founder DM)\n"
        "• Evidence-only; Brand DNA user-provided or generated, then locked for all images",
    )

    add_heading(doc, "4. Prompt contract", 1)
    add_body(
        doc,
        "System prompt = full stage prompts/*.md (fixed product instructions).\n"
        "User prompt = stage input payload only (varies per customer / brief).\n"
        "Output schema = schemas/*.schema.json (contract for the model reply).",
    )

    add_heading(doc, "5. Artifacts per run", 1)
    add_code_block(
        doc,
        "outputs/run_YYYYMMDD_HHMMSS/\n"
        "  00_normalized_raw.txt\n"
        "  00_ingest_manifest.json\n"
        "  00_brand_dna_lock.txt          (written when writers run)\n"
        "  01_raw.json\n"
        "  01_agent_1_input.json\n"
        "  02_company_lmo.json\n"
        "  02_company_handoff.json\n"
        "  03_company_content_matrix.json\n"
        "  04_posts/post_XX_infographic.json (+ .png)\n"
        "  04_posts/post_XX_carousel.json (+ slide_YY.png)\n"
        "  00_summary.json\n\n"
        "Caption (paste-ready) = post_text inside each post JSON.",
        size=9,
    )

    add_heading(doc, "6. Stages — System / User / Schema (full)", 1)

    for stage in STAGES:
        add_heading(doc, f"Stage {stage['num']} — {stage['title']}", 2)
        add_meta(doc, "Model / engine", stage["model"])
        add_meta(doc, "Notes", stage["notes"])

        add_heading(doc, "SYSTEM PROMPT", 3)
        if stage["prompt"]:
            prompt_path = PROMPTS / stage["prompt"]
            add_meta(doc, "File", f"prompts/{stage['prompt']}")
            add_body(
                doc,
                "Full stage prompt file (fixed product instructions). "
                "This entire block is sent as the system message.",
            )
            add_code_block(doc, prompt_path.read_text(encoding="utf-8"), size=8)
        else:
            add_body(doc, "No LLM system prompt (local ingest only).")

        add_heading(doc, "USER PROMPT", 3)
        if stage.get("user_prompt"):
            add_body(
                doc,
                "Pass only this stage’s input payload as the user message. "
                "Do not paste the system prompt here.",
            )
            add_code_block(doc, stage["user_prompt"], size=9)
        else:
            add_body(doc, "No separate user prompt template.")

        add_heading(doc, "OUTPUT SCHEMA", 3)
        if stage["schema"]:
            schema_path = SCHEMAS / stage["schema"]
            add_meta(doc, "File", f"schemas/{stage['schema']}")
            add_body(doc, "JSON Schema contract for the model reply / storage wrapper.")
            add_code_block(doc, schema_path.read_text(encoding="utf-8"), size=8)
        else:
            add_body(doc, "No JSON schema for this stage.")

        doc.add_paragraph()

    add_heading(doc, "7. Image render (no separate LLM prompt file)", 1)
    add_body(
        doc,
        "After 05c/06c return image_generation_prompt or per-slide image_gen_prompt, "
        "the runner builds 00_brand_dna_lock.txt from handoff.brand_dna, prepends it to "
        "each image prompt, and calls OpenAI gpt-image-2 (default size 1024x1536).",
    )

    add_heading(doc, "8. File index", 1)
    for stage in STAGES:
        if stage["prompt"] or stage["schema"]:
            sys_f = f"System=prompts/{stage['prompt']}" if stage["prompt"] else "System=n/a"
            user_f = {
                "01": "User=00_normalized_raw.txt",
                "02b": "User=01_agent_1_input.json",
                "03c": "User=02_company_handoff.json",
                "05c": "User=company_handoff + Infographic brief",
                "06c": "User=company_handoff + Carousel brief",
            }.get(stage["num"], "User=n/a")
            sch_f = f"Schema=schemas/{stage['schema']}" if stage["schema"] else "Schema=n/a"
            add_body(doc, f"Stage {stage['num']}: {sys_f} | {user_f} | {sch_f}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Wrote {out_path}")


def main() -> int:
    out = (
        Path(sys.argv[1])
        if len(sys.argv) > 1
        else Path.home() / "Downloads" / "B2B_company_page_PRD.docx"
    )
    build_doc(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
