"""Export all pipeline prompts + schemas into one structured Word doc for team sharing.

Production split (documented in the DOCX):
  System prompt  = full stage prompts/*.md (fixed product instructions)
  User prompt    = stage input payload only (varies per customer / brief)
  Output schema  = schemas/*.schema.json (contract for the model reply)
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]

STAGES = [
    {
        "num": "00",
        "title": "Ingest + Normalize",
        "model": "Local (scripts/ingest_normalize.py) — no LLM prompt",
        "prompt": None,
        "schema": None,
        "user_prompt": None,
        "notes": (
            "Accepts a legacy .txt file or a multi-upload pack. "
            "Writes 00_normalized_raw.txt and 00_ingest_manifest.json. "
            "Optional Gemini vision captions for images in packs."
        ),
    },
    {
        "num": "01",
        "title": "RAW Input Structurer",
        "model": "Gemini 3.5 Flash",
        "prompt": "01_raw_input_structurer.md",
        "schema": "01_raw_output.schema.json",
        "user_prompt": (
            "Normalized raw text from Stage 00 (contents of 00_normalized_raw.txt).\n\n"
            "This is the customer/onboarding paste: LinkedIn profile, company page, "
            "positioning axes, past posts, brand DNA, etc. — not the system prompt.\n\n"
            "Example shape:\n"
            "=== ONBOARDING / POSITIONING ===\n"
            "selected_positioning_mode: ...\n"
            "selected_growth_goal: ...\n"
            "...\n"
            "=== LINKEDIN PROFILE ===\n"
            "...\n"
            "=== COMPANY PAGE ===\n"
            "...\n"
            "=== PAST POSTS ===\n"
            "..."
        ),
        "notes": (
            "Output: 01_raw.json + 01_agent_1_input.json (payload for LMO)."
        ),
    },
    {
        "num": "02",
        "title": "LMO (LinkedIn Marketing Optimization)",
        "model": "Gemini 3.5 Flash",
        "prompt": "02_lmo.md",
        "schema": "02_lmo_output.schema.json",
        "user_prompt": (
            "agent_1_input JSON from Stage 01 (pretty-printed).\n\n"
            "This is the structured onboarding + profile + business payload only. "
            "Do not re-send the RAW system prompt here.\n\n"
            "Source artifact: 01_agent_1_input.json"
        ),
        "notes": (
            "Output: 02_lmo.json + 02_handoff.json (canonical payload for Matrix + writers). "
            "Produces headline, about, experience, pillars, Brand DNA, growth constraints."
        ),
    },
    {
        "num": "02b",
        "title": "Company Page LMO",
        "model": "gpt-6-luna (default). Rollback: TEXT_PROVIDER=gemini",
        "prompt": "02_company_lmo.md",
        "schema": "02_company_lmo_output.schema.json",
        "user_prompt": (
            "agent_1_input JSON from Stage 01 (pretty-printed).\n\n"
            "Same payload as personal LMO. Company-page track does not send the personal "
            "LMO system prompt.\n\n"
            "Source artifact: 01_agent_1_input.json\n\n"
            "Runner: python scripts/run_pipeline.py inputs/your_raw.txt --track company_page"
        ),
        "notes": (
            "Parallel Stage 02. Output: 02_company_lmo.json + 02_company_handoff.json. "
            "Official company_title passthrough, 5 taglines (max 120) + recommended, "
            "one About (1,000–1,500 characters), Brand DNA, company pillars, evidence bank. "
            "Feeds Company Matrix + company writers."
        ),
    },
    {
        "num": "03",
        "title": "Content Matrix",
        "model": "Gemini 3.5 Flash",
        "prompt": "03_content_matrix.md",
        "schema": "03_content_matrix_output.schema.json",
        "user_prompt": (
            "handoff JSON from Stage 02 (pretty-printed).\n\n"
            "Canonical LMO handoff only — confirmed niche, positioning, growth goal, "
            "pillars, Brand DNA, voice, credibility, content_mode_constraints, etc.\n\n"
            "Source artifact: 02_handoff.json"
        ),
        "notes": (
            "Output: 03_content_matrix.json with weekly_content_matrix "
            "(hard quotas: 3 Text + 2 Carousel + 3 Infographic)."
        ),
    },
    {
        "num": "03c",
        "title": "Company Content Matrix",
        "model": "gpt-6-luna (default). Rollback: TEXT_PROVIDER=gemini",
        "prompt": "03_company_content_matrix.md",
        "schema": "03_company_content_matrix_output.schema.json",
        "user_prompt": (
            "company handoff JSON from Stage 02b (pretty-printed).\n\n"
            "Source artifact: 02_company_handoff.json\n\n"
            "Runner: --track company_page"
        ),
        "notes": (
            "Exactly 4 briefs: 2 Infographic + 2 Carousel. "
            "company_post_type: educational_howto | case_study | behind_the_scenes | company_news. "
            "No polls. No Text."
        ),
    },
    {
        "num": "04",
        "title": "Text Post Writer",
        "model": "Gemini 3.5 Flash",
        "prompt": "04_text_post.md",
        "schema": "04_text_post_output.schema.json",
        "user_prompt": (
            "Built by build_writer_payload(handoff, post_brief) in scripts/run_pipeline.py:\n\n"
            "=== HANDOFF ===\n"
            "{ full 02_handoff.json }\n\n"
            "=== POST_BRIEF ===\n"
            "{ one Text Post object from weekly_content_matrix }\n\n"
            "User prompt = handoff + this brief only. System prompt = 04_text_post.md."
        ),
        "notes": (
            "Runner stores paste-ready post_text in 04_posts/post_XX_text.json."
        ),
    },
    {
        "num": "05",
        "title": "Infographic Writer + Image",
        "model": "Gemini 3.5 Flash (caption + image prompt) → OpenAI gpt-image-2 (PNG)",
        "prompt": "05_infographic_post.md",
        "schema": "05_infographic_output.schema.json",
        "user_prompt": (
            "Built by build_writer_payload(handoff, post_brief):\n\n"
            "=== HANDOFF ===\n"
            "{ full 02_handoff.json }\n\n"
            "=== POST_BRIEF ===\n"
            "{ one Infographic object from weekly_content_matrix }\n\n"
            "Gemini returns JSON (caption + image_generation_prompt). "
            "OpenAI then renders the PNG with Brand DNA lock prepended to the image prompt "
            "(image step is not this system/user pair)."
        ),
        "notes": (
            "OpenAI renders post_XX_infographic.png with Brand DNA lock."
        ),
    },
    {
        "num": "05c",
        "title": "Company Infographic Writer + Image",
        "model": "gpt-6-luna → OpenAI gpt-image-2",
        "prompt": "05_company_infographic_post.md",
        "schema": "05_infographic_output.schema.json",
        "user_prompt": (
            "Built by build_writer_payload(company_handoff, post_brief):\n\n"
            "=== HANDOFF ===\n"
            "{ full 02_company_handoff.json }\n\n"
            "=== POST_BRIEF ===\n"
            "{ one Infographic from 03_company_content_matrix.json }\n\n"
            "Company voice; company CTA; same Brand DNA image lock as founder track."
        ),
        "notes": (
            "Company-page fork of Stage 05. Reuses 05 schema + gpt-image-2."
        ),
    },
    {
        "num": "06",
        "title": "Carousel Writer + Slide Images",
        "model": "Gemini 3.5 Flash (caption + slides) → OpenAI gpt-image-2 (PNG per slide)",
        "prompt": "06_carousel_post.md",
        "schema": "06_carousel_output.schema.json",
        "user_prompt": (
            "Built by build_writer_payload(handoff, post_brief):\n\n"
            "=== HANDOFF ===\n"
            "{ full 02_handoff.json }\n\n"
            "=== POST_BRIEF ===\n"
            "{ one Carousel object from weekly_content_matrix }\n\n"
            "Gemini returns JSON (post_text + slides[]). "
            "OpenAI renders one PNG per slide with Brand DNA lock "
            "(image step is not this system/user pair)."
        ),
        "notes": (
            "OpenAI renders one PNG per slide (Brand DNA lock)."
        ),
    },
    {
        "num": "06c",
        "title": "Company Carousel Writer + Slide Images",
        "model": "gpt-6-luna → OpenAI gpt-image-2",
        "prompt": "06_company_carousel_post.md",
        "schema": "06_carousel_output.schema.json",
        "user_prompt": (
            "Built by build_writer_payload(company_handoff, post_brief):\n\n"
            "=== HANDOFF ===\n"
            "{ full 02_company_handoff.json }\n\n"
            "=== POST_BRIEF ===\n"
            "{ one Carousel from 03_company_content_matrix.json }\n\n"
            "Company voice; company CTA slide; same Brand DNA lock."
        ),
        "notes": (
            "Company-page fork of Stage 06. Reuses 06 schema + gpt-image-2."
        ),
    },
]


def set_run_font(run, *, size: int = 10, mono: bool = False, bold: bool = False) -> None:
    run.font.size = Pt(size)
    run.bold = bold
    if mono:
        run.font.name = "Consolas"
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:eastAsia"), "Consolas")
        rFonts.set(qn("w:ascii"), "Consolas")
        rFonts.set(qn("w:hAnsi"), "Consolas")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)


def add_meta(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    r.bold = True
    r.font.size = Pt(11)
    r2 = p.add_run(value)
    r2.font.size = Pt(11)


def add_body(doc: Document, text: str) -> None:
    for line in (text or "").splitlines() or [""]:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(2)
        for run in p.runs:
            run.font.size = Pt(11)


def add_code_block(doc: Document, text: str, *, size: int = 8) -> None:
    """Add large prompt/schema text as mono paragraphs (Word-friendly)."""
    for line in (text or "").splitlines() or [""]:
        safe = line.replace("\t", "    ")
        if len(safe) > 500:
            while safe:
                chunk, safe = safe[:500], safe[500:]
                p = doc.add_paragraph(chunk)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, size=size, mono=True)
        else:
            p = doc.add_paragraph(safe if safe else " ")
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            for run in p.runs:
                set_run_font(run, size=size, mono=True)


def build_doc(out_path: Path) -> None:
    prompts_dir = ROOT / "prompts"
    schemas_dir = ROOT / "schemas"

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    add_heading(doc, "B2B LinkedIn Growth OS — Full Process Pack", 0)
    add_body(
        doc,
        "Team reference: end-to-end pipeline with production System / User / Schema split. "
        "Source of truth remains prompts/ and schemas/ in the repo.",
    )
    doc.add_paragraph()

    add_heading(doc, "1. Pipeline sequence", 1)
    sequence = """Upload pack OR legacy .txt
  → 00 Ingest + Normalize            (local — no LLM system/user)
  → 01 RAW        System=01_raw…md          User=normalized raw text         Schema=01
  → 02 LMO        System=02_lmo.md          User=agent_1_input JSON          Schema=02
  → 03 Matrix     System=03_content…md      User=handoff JSON                Schema=03
  → 04 Text       System=04_text…md         User=handoff + post brief        Schema=04
  → 05 Infographic System=05_infographic…md User=handoff + post brief        Schema=05
       → 05b PNG  (OpenAI gpt-image-2 — image API, not this system/user pair)
  → 06 Carousel   System=06_carousel…md     User=handoff + post brief        Schema=06
       → 06b PNGs (OpenAI gpt-image-2 — one call per slide)

Each Gemini stage: System prompt (fixed .md) + User prompt (payload only) → model output validated by schema.

Runner: scripts/run_pipeline.py
Hard week mix: 3 Text + 2 Carousel + 3 Infographic
Axes: selected_positioning_mode × selected_growth_goal"""
    add_code_block(doc, sequence, size=9)

    doc.add_paragraph()
    add_heading(doc, "2. System vs User vs Schema", 1)
    add_body(
        doc,
        "SYSTEM PROMPT — Full stage prompt file (prompts/0X_….md). Fixed product instructions. "
        "In production this goes into the model’s system / systemInstruction channel. "
        "Does not change per customer.",
    )
    add_body(
        doc,
        "USER PROMPT — Stage input payload only. Varies per run: raw text, agent_1_input, "
        "handoff, or handoff + one Matrix brief. Never duplicate the system prompt here.",
    )
    add_body(
        doc,
        "OUTPUT SCHEMA — schemas/0X_….schema.json. Formal contract for the model’s reply "
        "(required keys, types). Blueprint, not customer content. Not pasted into system/user "
        "unless you optionally add a short “match schema” reminder.",
    )
    add_body(
        doc,
        "Local note: scripts/run_pipeline.py call_gemini() currently concatenates system + "
        "===== USER PAYLOAD ===== + input into one Gemini user message. Production should map "
        ".md → systemInstruction/system and payload → user message for cleaner caching and boundaries.",
    )

    doc.add_page_break()
    add_heading(doc, "3. Stages — System prompt, User prompt, Output schema", 1)

    for stage in STAGES:
        add_heading(doc, f"Stage {stage['num']} — {stage['title']}", 1)
        add_meta(doc, "Model / engine", stage["model"])
        add_meta(doc, "Notes", stage["notes"])

        # System prompt
        add_heading(doc, "A. SYSTEM PROMPT (production system channel)", 2)
        if stage["prompt"]:
            prompt_path = prompts_dir / stage["prompt"]
            add_meta(doc, "File", f"prompts/{stage['prompt']}")
            add_body(doc, "Paste the entire file below into the production system prompt.")
            add_code_block(doc, prompt_path.read_text(encoding="utf-8"), size=8)
        else:
            add_body(doc, "No LLM system prompt (local ingest only).")

        # User prompt
        add_heading(doc, "B. USER PROMPT (production user channel — input only)", 2)
        if stage.get("user_prompt"):
            add_body(
                doc,
                "Pass only this stage’s input payload as the user message. "
                "Do not prepend the system prompt again.",
            )
            add_code_block(doc, stage["user_prompt"], size=9)
        else:
            add_body(doc, "No LLM user prompt (local ingest only).")

        # Schema
        add_heading(doc, "C. OUTPUT SCHEMA (reply contract)", 2)
        if stage["schema"]:
            schema_path = schemas_dir / stage["schema"]
            add_meta(doc, "File", f"schemas/{stage['schema']}")
            add_code_block(doc, schema_path.read_text(encoding="utf-8"), size=8)
        else:
            add_body(doc, "No JSON schema for this stage.")

        doc.add_page_break()

    add_heading(doc, "4. How to run", 1)
    add_code_block(
        doc,
        """pip install -r requirements.txt

# .env
# GEMINI_API_KEY=...
# GEMINI_MODEL=gemini-3.5-flash
# OPENAI_API_KEY=...
# OPENAI_IMAGE_MODEL=gpt-image-2

python scripts/run_pipeline.py inputs/your_raw.txt
python scripts/run_pipeline.py inputs/your_raw.txt --stop-after lmo
python scripts/run_pipeline.py --resume outputs/run_YYYYMMDD_HHMMSS

python scripts/export_content_pack_docx.py run_YYYYMMDD_HHMMSS path/to/Downloads/pack.docx
python scripts/export_prompts_schemas_docx.py path/to/Downloads/B2B_pipeline_prompts_and_schemas.docx
""",
        size=9,
    )

    add_heading(doc, "5. File index (System / User / Schema)", 1)
    for stage in STAGES:
        if stage["prompt"] or stage["schema"]:
            sys_f = f"System=prompts/{stage['prompt']}" if stage["prompt"] else "System=n/a"
            user_f = {
                "01": "User=00_normalized_raw.txt",
                "02": "User=01_agent_1_input.json",
                "02b": "User=01_agent_1_input.json",
                "03": "User=02_handoff.json",
                "03c": "User=02_company_handoff.json",
                "04": "User=handoff + Text brief",
                "05": "User=handoff + Infographic brief",
                "05c": "User=company_handoff + Infographic brief",
                "06": "User=handoff + Carousel brief",
                "06c": "User=company_handoff + Carousel brief",
            }.get(stage["num"], "User=n/a")
            sch_f = f"Schema=schemas/{stage['schema']}" if stage["schema"] else "Schema=n/a"
            add_body(doc, f"Stage {stage['num']}: {sys_f} | {user_f} | {sch_f}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"Wrote {out_path}")


def main() -> int:
    out = (
        Path(sys.argv[1])
        if len(sys.argv) > 1
        else Path.home() / "Downloads" / "B2B_pipeline_prompts_and_schemas.docx"
    )
    build_doc(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
